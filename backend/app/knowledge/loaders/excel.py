from langchain_core.documents import Document
from openpyxl import load_workbook


def load(file_path: str) -> list[Document]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                parts.append(f"# {sheet.title}\n" + "\n".join(rows))
    finally:
        workbook.close()
    return [Document(page_content="\n\n".join(parts), metadata={"source": file_path})]
