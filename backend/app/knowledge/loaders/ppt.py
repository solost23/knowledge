from langchain_core.documents import Document
from pptx import Presentation


def load(file_path: str) -> list[Document]:
    parts = []
    for number, slide in enumerate(Presentation(file_path).slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        if lines:
            parts.append(f"# 幻灯片 {number}\n" + "\n".join(lines))
    if not parts:
        return []
    return [Document(page_content="\n\n".join(parts), metadata={"source": file_path})]
